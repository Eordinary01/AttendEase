import re
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

def set_heading_style(run, size, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, italic=False):
    font = run.font
    font.size = Pt(size)
    font.bold = bold
    font.italic = italic
    font.color.rgb = RGBColor(0, 0, 0)
    run.alignment = align

def add_heading_paragraph(doc, text, level=1):
    if level == 1:
        size = 18
    elif level == 2:
        size = 16
    elif level == 3:
        size = 14
    else:
        size = 12
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_heading_style(run, size, bold=True)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    return p

def add_normal_paragraph(doc, text, italic=False, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = alignment
    run = p.add_run(text)
    set_heading_style(run, 12, bold=False, italic=italic)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    return p

def add_code_block(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_heading_style(run, 10, bold=False)
    run.font.name = "Consolas"
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return p

def parse_md_table(lines):
    # Simple markdown table parser -> list of rows
    rows = []
    for line in lines:
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(set(c.replace("-", "").replace(":", "")) == set() for c in cells):
            continue
        rows.append(cells)
    return rows

def insert_table(doc, rows):
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = cell_text
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    set_heading_style(run, 11, bold=(i == 0))
    doc.add_paragraph()

def md_to_docx(md_path, docx_path):
    with open(md_path, "r", encoding="utf-8") as f:
        md_text = f.read()
    doc = Document()
    # Set default font for document
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(12)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    style.paragraph_format.space_after = Pt(6)

    lines = md_text.splitlines()
    i = 0
    in_code = False
    code_lines = []
    table_lines = []
    in_table = False

    def flush_table():
        nonlocal table_lines, in_table
        if table_lines:
            insert_table(doc, parse_md_table(table_lines))
            table_lines = []
        in_table = False

    def flush_code():
        nonlocal code_lines, in_code
        if code_lines:
            add_code_block(doc, "\n".join(code_lines))
            code_lines = []
        in_code = False

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                flush_code()
            else:
                flush_table()
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if stripped.startswith("|"):
            in_table = True
            table_lines.append(stripped)
            i += 1
            continue
        elif in_table:
            flush_table()

        if stripped.startswith("# "):
            flush_table()
            add_heading_paragraph(doc, stripped.lstrip("# ").strip(), level=1)
        elif stripped.startswith("## "):
            flush_table()
            add_heading_paragraph(doc, stripped.lstrip("# ").strip(), level=2)
        elif stripped.startswith("### "):
            flush_table()
            add_heading_paragraph(doc, stripped.lstrip("# ").strip(), level=3)
        elif stripped.startswith("#### "):
            flush_table()
            add_heading_paragraph(doc, stripped.lstrip("# ").strip(), level=4)
        elif stripped.startswith("- ") or re.match(r"^\d+\.\s", stripped):
            text = re.sub(r"^[- ]\s*", "", stripped)
            text = re.sub(r"^\d+\.\s*", "", text)
            text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(text)
            set_heading_style(run, 12, bold=False)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            p.paragraph_format.space_after = Pt(3)
        elif stripped == "":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
        else:
            text = stripped
            text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
            text = re.sub(r"\*(.*?)\*", r"\1", text)
            if text.startswith("_") and text.endswith("_"):
                text = text.strip("_")
                add_normal_paragraph(doc, text, italic=True)
            else:
                add_normal_paragraph(doc, text)
        i += 1

    flush_table()
    flush_code()
    doc.save(docx_path)
    print("Generated DOCX:", docx_path, "size:", os.path.getsize(docx_path), "bytes")

md_files = [
    r"C:\Users\Parth Manocha\Desktop\DEV\attendease_\docs\Synopsis_BTech_AttendEase.md",
    r"C:\Users\Parth Manocha\Desktop\DEV\attendease_\docs\Report_BTech_AttendEase.md",
]
for md_path in md_files:
    docx_path = md_path.replace(".md", ".docx")
    md_to_docx(md_path, docx_path)
