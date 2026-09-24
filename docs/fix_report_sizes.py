from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
import os, re

SRC = r"C:\Users\Parth Manocha\Desktop\DEV\attendease_\docs"
REF = r"C:\Users\Parth Manocha\Downloads"

ref_syn = os.path.join(REF, "Onect_Synopsis_Refactored (1).docx")
ref_rpt = os.path.join(REF, "Onect_Report.docx")

src_rpt_md = os.path.join(SRC, "Report_BTech_AttendEase.md")
out_rpt = os.path.join(SRC, "Report_BTech_AttendEase_Layout.docx")

logo_path = r"C:\Users\Parth Manocha\Desktop\DEV\attendease_\docs\ref_media\Onect_Report_image1.png"

ref_rpt_doc = Document(ref_rpt)

def set_run(run, size_pt=None, bold=None, italic=None, name="Times New Roman"):
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if name:
        run.font.name = name

def add_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.LEFT, size_pt=12, bold=False, italic=False, space_after=6, space_before=0, keep_with_next=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if keep_with_next:
        p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run(run, size_pt=size_pt, bold=bold, italic=italic)
    return p

def add_heading(doc, text, level, h1=18, h2=16, h3=14, body=12):
    if level == 1:
        size_pt = h1
        bold = True
    elif level == 2:
        size_pt = h2
        bold = True
    elif level == 3:
        size_pt = h3
        bold = True
    else:
        size_pt = body
        bold = False
    add_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.LEFT, size_pt=size_pt, bold=bold, space_after=6, space_before=12, keep_with_next=True)

def set_header_footer(section, header_text, footer_text, header_size=10, footer_size=10):
    header = section.header
    header.is_linked_to_previous = False
    for p in header.paragraphs:
        p.clear()
    if header_text:
        p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(header_text)
        set_run(run, size_pt=header_size, bold=True)
    footer = section.footer
    footer.is_linked_to_previous = False
    for p in footer.paragraphs:
        p.clear()
    if footer_text:
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(footer_text)
        set_run(run, size_pt=footer_size, bold=False)

def add_page_break(doc):
    doc.add_page_break()

def parse_md_table(lines):
    rows = []
    for line in lines:
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(set(c.replace("-", "").replace(":", "")) == set() for c in cells):
            continue
        rows.append(cells)
    return rows

def add_table(doc, rows, body_size=12):
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = cell_text
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    set_run(run, size_pt=body_size, bold=(i==0))
    doc.add_paragraph()

# Build report with 12pt body text matching reference report
with open(src_rpt_md, "r", encoding="utf-8") as f:
    rpt_md = f.read()

rpt_doc = Document()
rpt_doc.sections[0].page_height = ref_rpt_doc.sections[0].page_height
rpt_doc.sections[0].page_width = ref_rpt_doc.sections[0].page_width
rpt_doc.sections[0].top_margin = ref_rpt_doc.sections[0].top_margin
rpt_doc.sections[0].bottom_margin = ref_rpt_doc.sections[0].bottom_margin
rpt_doc.sections[0].left_margin = ref_rpt_doc.sections[0].left_margin
rpt_doc.sections[0].right_margin = ref_rpt_doc.sections[0].right_margin
rpt_doc.sections[0].different_first_page_header_footer = True

# Title page
add_paragraph(rpt_doc, "MAJOR PROJECT REPORT", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=14, bold=True, space_after=12)
add_paragraph(rpt_doc, "AttendEase", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=28, bold=True, space_after=6)
add_paragraph(rpt_doc, "A Multi-Tenant SaaS Academic ERP with Intelligent Attendance, Exam & Fee Management", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=16, bold=False, space_after=12)
add_paragraph(rpt_doc, "Submitted in partial fulfillment of the requirements for the award of the degree of", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=False, space_after=6)
add_paragraph(rpt_doc, "BACHELOR OF TECHNOLOGY", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=14, bold=True, space_after=6)
add_paragraph(rpt_doc, "in", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=False, space_after=6)
add_paragraph(rpt_doc, "COMPUTER SCIENCE AND ENGINEERING", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=14, bold=True, space_after=24)
add_paragraph(rpt_doc, "Submitted By:", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=True, space_after=6)
add_paragraph(rpt_doc, "Parth Manocha (23BCON0051)", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=False, space_after=24)
add_paragraph(rpt_doc, "Under the Supervision of", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=True, space_after=6)
add_paragraph(rpt_doc, "Mr. Divyansh Dewani", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=True, space_after=6)
add_paragraph(rpt_doc, "Teaching Assistant", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=False, space_after=6)
add_paragraph(rpt_doc, "Department of Computer Science and Engineering", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=False, space_after=24)
add_paragraph(rpt_doc, "JECRC University, Jaipur", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=True, space_after=12)
add_paragraph(rpt_doc, "2025-26", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=False, space_after=24)

if os.path.exists(logo_path):
    rpt_doc.add_picture(logo_path, width=Inches(1.5))
    last_paragraph = rpt_doc.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

add_page_break(rpt_doc)

new_section = rpt_doc.add_section()
new_section.page_height = ref_rpt_doc.sections[0].page_height
new_section.page_width = ref_rpt_doc.sections[0].page_width
new_section.top_margin = ref_rpt_doc.sections[0].top_margin
new_section.bottom_margin = ref_rpt_doc.sections[0].bottom_margin
new_section.left_margin = ref_rpt_doc.sections[0].left_margin
new_section.right_margin = ref_rpt_doc.sections[0].right_margin
set_header_footer(new_section, "AttendEase – Multi-Tenant SaaS Academic ERP", "2025-26\t\t1", header_size=10, footer_size=10)

# Parse MD and add content
lines = rpt_md.splitlines()
i = 0
in_code = False
code_lines = []
table_lines = []
in_table = False

def flush_table(doc):
    global table_lines, in_table
    if table_lines:
        add_table(doc, parse_md_table(table_lines), body_size=12)
        table_lines = []
    in_table = False

def flush_code(doc):
    global code_lines, in_code
    if code_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run("\n".join(code_lines))
        set_run(run, size_pt=10, bold=False, name="Consolas")
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        code_lines = []
    in_code = False

while i < len(lines):
    line = lines[i]
    stripped = line.strip()

    if stripped.startswith("```"):
        if in_code:
            flush_code(rpt_doc)
        else:
            flush_table(rpt_doc)
            in_code = True
        i += 1
        continue

    if in_code:
        code_lines.append(line)
        i += 1
        continue

    if stripped.startswith("|"):
        in_table = True
        table_lines.append(stripped)
        i += 1
        continue
    elif in_table:
        flush_table(rpt_doc)

    if stripped.startswith("# "):
        flush_table(rpt_doc)
        add_heading(rpt_doc, stripped.lstrip("# ").strip(), 1, h1=18, h2=16, h3=14, body=12)
    elif stripped.startswith("## "):
        flush_table(rpt_doc)
        add_heading(rpt_doc, stripped.lstrip("# ").strip(), 2, h1=18, h2=16, h3=14, body=12)
    elif stripped.startswith("### "):
        flush_table(rpt_doc)
        add_heading(rpt_doc, stripped.lstrip("# ").strip(), 3, h1=18, h2=16, h3=14, body=12)
    elif stripped.startswith("#### "):
        flush_table(rpt_doc)
        add_heading(rpt_doc, stripped.lstrip("# ").strip(), 4, h1=18, h2=16, h3=14, body=12)
    elif stripped.startswith("- ") or re.match(r"^\d+\.\s", stripped):
        text = re.sub(r"^[- ]\s*", "", stripped)
        text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
        text = re.sub(r"\*(.*?)\*", r"\1", text)
        p = rpt_doc.add_paragraph(style="List Bullet")
        run = p.add_run(text)
        set_run(run, size_pt=12, bold=False)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.space_after = Pt(3)
    elif stripped == "":
        p = rpt_doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
    else:
        text = stripped
        text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
        text = re.sub(r"\*(.*?)\*", r"\1", text)
        if text.startswith("_") and text.endswith("_"):
            text = text.strip("_")
            add_paragraph(rpt_doc, text, italic=True, size_pt=12, space_after=6)
        else:
            add_paragraph(rpt_doc, text, size_pt=12, space_after=6)
    i += 1

flush_table(rpt_doc)
flush_code(rpt_doc)
rpt_doc.save(out_rpt)
print("Generated:", out_rpt)
